from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable


SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx"}


def _cell_text(value: object) -> str:
    """Turn spreadsheet/document values into compact, Markdown-safe text."""
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").replace("|", "\\|").replace("\n", "<br>")


def _markdown_table(rows: Iterable[Iterable[object]]) -> str:
    normalized = [[_cell_text(value) for value in row] for row in rows]
    normalized = [row for row in normalized if any(cell for cell in row)]
    if not normalized:
        return ""
    width = max(len(row) for row in normalized)
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    lines = ["| " + " | ".join(normalized[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _parse_pdf(source: Path) -> str:
    """Extract the text layer from a PDF; deliberately does not run OCR."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(source))
    pages: list[str] = []
    try:
        for page_number in range(len(document)):
            page = document[page_number]
            try:
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_bounded().strip()
                finally:
                    text_page.close()
            finally:
                page.close()
            if text:
                pages.append(f"## 第 {page_number + 1} 页\n\n{text}")
    finally:
        document.close()
    return "\n\n".join(pages)


def _parse_docx(source: Path) -> str:
    from docx import Document
    from docx.table import Table

    document = Document(str(source))
    blocks: list[str] = []
    content = document.iter_inner_content() if hasattr(document, "iter_inner_content") else [*document.paragraphs, *document.tables]
    for block in content:
        if isinstance(block, Table):
            table = _markdown_table([[cell.text for cell in row.cells] for row in block.rows])
            if table:
                blocks.append(table)
            continue
        text = block.text.strip()
        if not text:
            continue
        style_name = getattr(getattr(block, "style", None), "name", "") or ""
        if style_name.lower().startswith("heading"):
            try:
                level = int(style_name.rsplit(" ", 1)[-1])
            except ValueError:
                level = 1
            blocks.append(f"{'#' * max(1, min(level, 6))} {text}")
        elif "list" in style_name.lower():
            blocks.append(f"- {text}")
        else:
            blocks.append(text)
    return "\n\n".join(blocks)


def _parse_xlsx(source: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(source), read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for sheet in workbook.worksheets:
            table = _markdown_table(sheet.iter_rows(values_only=True))
            if table:
                sections.append(f"## {sheet.title}\n\n{table}")
    finally:
        workbook.close()
    return "\n\n".join(sections)


def _iter_ppt_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        nested = getattr(shape, "shapes", None)
        if nested is not None:
            yield from _iter_ppt_shapes(nested)
        else:
            yield shape


def _parse_pptx(source: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(source))
    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in _iter_ppt_shapes(slide.shapes):
            if getattr(shape, "has_table", False):
                table = _markdown_table([[cell.text for cell in row.cells] for row in shape.table.rows])
                if table:
                    blocks.append(table)
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    blocks.append(text)
        if blocks:
            slides.append(f"## 第 {slide_number} 页\n\n" + "\n\n".join(blocks))
    return "\n\n".join(slides)


def parse_document(path: str, filename: str | None = None) -> dict[str, Any]:
    """Parse modern PDF/Office documents with small, format-specific readers."""
    source = Path(path)
    extension = source.suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise ValueError("仅支持 PDF、DOCX、XLSX、PPTX；不支持旧版 DOC、XLS、PPT")
    if not source.is_file():
        raise FileNotFoundError(f"文档不存在：{source}")

    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".pptx": _parse_pptx,
    }
    markdown = parsers[extension](source)
    return {
        "filename": filename or source.name,
        "extension": extension,
        "markdown": markdown,
        "charCount": len(markdown),
    }
