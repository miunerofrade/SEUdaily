from __future__ import annotations

from pathlib import Path

import pytest

from cvstream.document_parser import parse_document


def _write_text_pdf(path: Path, text: str) -> None:
    content = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(output)


def test_parse_pdf_without_ocr_or_model_download(tmp_path: Path) -> None:
    source = tmp_path / "sample.pdf"
    _write_text_pdf(source, "PDF parser works")

    result = parse_document(str(source))

    assert result["extension"] == ".pdf"
    assert "PDF parser works" in result["markdown"]


def test_parse_modern_office_formats(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("DOCX parser works")
    document.save(docx_path)

    xlsx_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "XLSX parser works"
    workbook.save(xlsx_path)

    pptx_path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX parser works"
    presentation.save(pptx_path)

    assert "DOCX parser works" in parse_document(str(docx_path))["markdown"]
    assert "XLSX parser works" in parse_document(str(xlsx_path))["markdown"]
    assert "PPTX parser works" in parse_document(str(pptx_path))["markdown"]


def test_rejects_legacy_office_formats(tmp_path: Path) -> None:
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"not a modern office document")

    with pytest.raises(ValueError, match="不支持旧版"):
        parse_document(str(source))
